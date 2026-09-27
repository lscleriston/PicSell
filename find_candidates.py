# -*- coding: utf-8 -*-
import json, re

with open('dados_produtos.json', 'r', encoding='utf-8') as f:
    loja_data = json.load(f)

loja_items = loja_data['itens_detalhados']

# Map of manual/exact overrides for items without SKU or where titles differ
# Let's inspect each SEM SKU item and find the candidate in loja_items
import openpyxl, io, zipfile

def carregar_shopee_workbook(caminho):
    with zipfile.ZipFile(caminho, 'r') as z_in:
        out_buf = io.BytesIO()
        with zipfile.ZipFile(out_buf, 'w') as z_out:
            for item in z_in.infolist():
                data = z_in.read(item.filename)
                if item.filename.endswith('sheet1.xml'):
                    data = data.replace(b'activePane="bottom_left"', b'activePane="bottomLeft"')
                z_out.writestr(item, data)
    out_buf.seek(0)
    return openpyxl.load_workbook(out_buf, data_only=True)

wb = carregar_shopee_workbook('mass_update_sales_info_1634195952_20260927054221.xlsx')
sheet = wb.active

sem_sku = {}
for r in range(7, sheet.max_row + 1):
    pid = sheet.cell(row=r, column=1).value
    pname = str(sheet.cell(row=r, column=2).value or '').strip()
    vname = str(sheet.cell(row=r, column=4).value or '').strip()
    psku = str(sheet.cell(row=r, column=5).value or '').strip()
    vsku = str(sheet.cell(row=r, column=6).value or '').strip()
    price = sheet.cell(row=r, column=7).value
    if pid and not psku and not vsku:
        if pid not in sem_sku:
            sem_sku[pid] = {'name': pname, 'price': price, 'vname': vname}

print(f"Total produtos sem SKU na Shopee: {len(sem_sku)}")

for pid, d in sem_sku.items():
    name = d['name'].lower()
    # Find matching candidates in loja_items
    candidates = []
    # Tokenize name
    words = [w for w in re.findall(r'\w+', name) if len(w) > 3 and w not in ['bebe', 'bebê', 'infantil', 'menino', 'menina', 'unissex', 'roupa']]
    for it in loja_items:
        it_name = it['nome'].lower()
        score = sum(1 for w in words if w in it_name)
        if score > 0:
            candidates.append((score, it))
    candidates.sort(key=lambda x: x[0], reverse=True)
    top = candidates[:3]
    print(f"\nSHOPEE ({pid}) [P: {d['price']}]: {d['name'][:50]}")
    for sc, it in top:
        print(f"   -> Score {sc}: {it['nome']} (SKU: {it['sku']} | Custo: {it['custo']} | Preco Loja: {it['preco_venda']})")
