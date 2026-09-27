# -*- coding: utf-8 -*-
import openpyxl, io, zipfile, json

def carregar_shopee_workbook(caminho):
    with zipfile.ZipFile(caminho, 'r') as z_in:
        out_buf = io.BytesIO()
        with zipfile.ZipFile(out_buf, 'w') as z_out:
            for item in z_in.infolist():
                data = z_in.read(item.filename)
                if item.filename.endswith('sheet1.xml'):
                    data = data.replace(b'activePane="bottom_left"', b'activePane="bottomLeft"')
                z_out.writestr(item, data)
    out_buf.seek(0)
    return openpyxl.load_workbook(out_buf, data_only=True)

wb = carregar_shopee_workbook('mass_update_sales_info_1634195952_20260927054221.xlsx')
sheet = wb.active
distinct_prods = {}
for r in range(7, sheet.max_row + 1):
    pid = sheet.cell(row=r, column=1).value
    pname = str(sheet.cell(row=r, column=2).value or '').strip()
    psku = str(sheet.cell(row=r, column=5).value or '').strip()
    price = sheet.cell(row=r, column=7).value
    if pid:
        if pid not in distinct_prods:
            distinct_prods[pid] = {'name': pname, 'sku': psku, 'price': price, 'count': 0}
        distinct_prods[pid]['count'] += 1

print(f"Total distinct products on Shopee: {len(distinct_prods)}")
for pid, d in distinct_prods.items():
    print(f"ID: {pid} | Sku: {d['sku']:15} | Preco: {d['price']} | Vars: {d['count']:2} | {d['name'][:60]}")
