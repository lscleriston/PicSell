# -*- coding: utf-8 -*-
import os
import glob
import openpyxl, io, zipfile, json, re

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SHOPEE_DIR = os.path.join(ROOT, 'data', 'raw', 'shopee')
PRODUTOS_JSON = os.path.join(ROOT, 'data', 'processed', 'dados_produtos.json')


def planilha_legada():
    encontrados = glob.glob(os.path.join(SHOPEE_DIR, 'mass_update_sales_info*.xlsx'))
    if not encontrados:
        raise SystemExit(f'Nenhuma mass_update_sales_info*.xlsx em {SHOPEE_DIR}')
    encontrados.sort(key=os.path.getmtime, reverse=True)
    return encontrados[0]

def carregar_shopee_workbook(caminho):
    with zipfile.ZipFile(caminho, 'r') as z_in:
        out_buf = io.BytesIO()
        with zipfile.ZipFile(out_buf, 'w') as z_out:
            for item in z_in.infolist():
                data = z_in.read(item.filename)
                if item.filename.endswith('sheet1.xml'):
                    data = data.replace(b'activePane="bottom_left"', b'activePane="bottomLeft"')
                z_out.writestr(item, data)
    out_buf.seek(0)
    return openpyxl.load_workbook(out_buf, data_only=True)

with open(PRODUTOS_JSON, 'r', encoding='utf-8') as f:
    loja_data = json.load(f)

loja_items = loja_data['itens_detalhados']
loja_by_sku = {str(it['sku']).strip().upper(): it for it in loja_items}

wb = carregar_shopee_workbook(planilha_legada())
sheet = wb.active

prods = {}
for r in range(7, sheet.max_row + 1):
    pid = sheet.cell(row=r, column=1).value
    pname = str(sheet.cell(row=r, column=2).value or '').strip()
    var_id = sheet.cell(row=r, column=3).value
    var_name = str(sheet.cell(row=r, column=4).value or '').strip()
    psku = str(sheet.cell(row=r, column=5).value or '').strip()
    vsku = str(sheet.cell(row=r, column=6).value or '').strip()
    price = sheet.cell(row=r, column=7).value
    stock = sheet.cell(row=r, column=9).value
    if pid:
        if pid not in prods:
            prods[pid] = {'name': pname, 'psku': psku, 'price': price, 'vars': []}
        prods[pid]['vars'].append({'var_id': var_id, 'var_name': var_name, 'vsku': vsku, 'price': price, 'stock': stock})

print(f"Total produtos Shopee: {len(prods)}")
for pid, d in list(prods.items()):
    pname = d['name']
    psku = d['psku']
    v_example = d['vars'][0]
    vsku = v_example['vsku']
    vname = v_example['var_name']
    
    # Check if direct SKU
    matched = None
    if vsku and vsku.upper() in loja_by_sku:
        matched = ('VSKU', loja_by_sku[vsku.upper()])
    elif psku and psku.upper() in loja_by_sku:
        matched = ('PSKU', loja_by_sku[psku.upper()])
    else:
        # Check partial
        for k, v in loja_by_sku.items():
            if (vsku and (k in vsku.upper() or vsku.upper() in k)) or (psku and (k in psku.upper() or psku.upper() in k)):
                matched = ('PARTIAL_SKU', v)
                break
    
    if matched:
        print(f"[OK-SKU {matched[0]}] {pname[:30]} -> {matched[1]['nome'][:30]} (Custo: {matched[1]['custo']})")
    else:
        print(f"[SEM SKU] {pname[:35]} | Var: {vname[:20]} | Shopee P: {d['price']}")
