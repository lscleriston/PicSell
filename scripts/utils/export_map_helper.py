# -*- coding: utf-8 -*-
import os
import glob
import openpyxl, io, zipfile, json

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SHOPEE_DIR = os.path.join(ROOT, 'data', 'raw', 'shopee')
PRODUTOS_JSON = os.path.join(ROOT, 'data', 'processed', 'dados_produtos.json')
MAP_HELPER = os.path.join(ROOT, 'scripts', 'config', 'map_helper.txt')


def planilha_legada():
    encontrados = glob.glob(os.path.join(SHOPEE_DIR, 'mass_update_sales_info*.xlsx'))
    if not encontrados:
        raise SystemExit(f'Nenhuma mass_update_sales_info*.xlsx em {SHOPEE_DIR}')
    encontrados.sort(key=os.path.getmtime, reverse=True)
    return encontrados[0]


def carregar_shopee_workbook(caminho):
    with zipfile.ZipFile(caminho, 'r') as z_in:
        out_buf = io.BytesIO()
        with zipfile.ZipFile(out_buf, 'w') as z_out:
            for item in z_in.infolist():
                data = z_in.read(item.filename)
                if item.filename.endswith('sheet1.xml'):
                    data = data.replace(b'activePane="bottom_left"', b'activePane="bottomLeft"')
                z_out.writestr(item, data)
    out_buf.seek(0)
    return openpyxl.load_workbook(out_buf, data_only=True)

with open(PRODUTOS_JSON, 'r', encoding='utf-8') as f:
    loja_data = json.load(f)

wb = carregar_shopee_workbook(planilha_legada())
sheet = wb.active

prods = {}
for r in range(7, sheet.max_row + 1):
    pid = sheet.cell(row=r, column=1).value
    pname = str(sheet.cell(row=r, column=2).value or '').strip()
    var_id = sheet.cell(row=r, column=3).value
    var_name = str(sheet.cell(row=r, column=4).value or '').strip()
    psku = str(sheet.cell(row=r, column=5).value or '').strip()
    vsku = str(sheet.cell(row=r, column=6).value or '').strip()
    price = sheet.cell(row=r, column=7).value
    if pid and pid not in prods:
        prods[pid] = {'pname': pname, 'psku': psku, 'vsku': vsku, 'price': price, 'vname': var_name}

print(f"Total produtos únicos Shopee: {len(prods)}")

# List all products in loja_data to make exact mapping easy
nomes_loja = sorted(list({it['nome'] for it in loja_data['itens_detalhados']}))

with open(MAP_HELPER, 'w', encoding='utf-8') as f:
    f.write("PRODUTOS SHOPEE:\n")
    for pid, d in prods.items():
        f.write(f"ID: {pid} | Sku: {d['psku']} | Price: {d['price']} | {d['pname']} | Var: {d['vname']}\n")
    f.write("\n\nPRODUTOS LOJA INTEGRADA:\n")
    for n in nomes_loja:
        f.write(f"{n}\n")

print("map_helper.txt gerado com sucesso!")
