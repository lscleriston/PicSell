# -*- coding: utf-8 -*-
"""
Módulo para ingestão e sincronização da planilha de catálogo da loja no SQLite.
Lê data/raw/catalogo/*.xlsx e alimenta as tabelas produtos_pai e produtos_variacao.
"""

import os
import sys
import glob
import re
import html
import json
import openpyxl

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass
from .database import get_connection, init_db, ROOT

PASTA_CATALOGO = os.path.join(ROOT, 'data', 'raw', 'catalogo')


def clean_html(raw_html):
    if not raw_html:
        return ''
    cleanr = re.compile(r'<.*?>')
    cleantext = re.sub(cleanr, ' ', str(raw_html))
    return ' '.join(html.unescape(cleantext).split()).strip()


def importar_catalogo(caminho_arquivo=None, db_path=None):
    """Importa o catálogo da loja para o banco de dados SQLite."""
    init_db(db_path)
    
    if not caminho_arquivo:
        arquivos = glob.glob(os.path.join(PASTA_CATALOGO, '*.xlsx'))
        if not arquivos:
            print(f"Erro: Nenhuma planilha .xlsx encontrada em {os.path.relpath(PASTA_CATALOGO, ROOT)}.")
            return False
        arquivos.sort(key=os.path.getmtime, reverse=True)
        caminho_arquivo = arquivos[0]

    nome_arquivo = os.path.basename(caminho_arquivo)
    print(f"📦 Importando catálogo: {nome_arquivo}...")

    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    sheet = wb.active

    conn = get_connection(db_path)
    cursor = conn.cursor()

    # 1. Mapeamento dos produtos pai (com-variacao)
    parents = {}
    for r in range(3, sheet.max_row + 1):
        tipo = sheet.cell(row=r, column=2).value
        if tipo == 'com-variacao':
            sku = str(sheet.cell(row=r, column=4).value or '').strip()
            cat1 = str(sheet.cell(row=r, column=32).value or '').strip()
            cat2 = str(sheet.cell(row=r, column=33).value or '').strip()
            categoria = cat1 + (' > ' + cat2 if cat2 else '')
            parents[sku] = {
                'sku_pai': sku,
                'nome': str(sheet.cell(row=r, column=11).value or '').strip(),
                'ativo': str(sheet.cell(row=r, column=5).value or 'S'),
                'destaque': str(sheet.cell(row=r, column=7).value or 'N'),
                'ncm': str(sheet.cell(row=r, column=8).value or '').strip(),
                'categoria': categoria,
                'cat1': cat1,
                'cat2': cat2,
                'imagem_url': str(sheet.cell(row=r, column=37).value or '').strip(),
                'imagem_url2': str(sheet.cell(row=r, column=38).value or '').strip(),
                'descricao': clean_html(sheet.cell(row=r, column=14).value)[:500],
                'variacoes': []
            }

    # 2. Processamento das variações e itens individuais
    items = []
    total_linhas = 0

    for r in range(3, sheet.max_row + 1):
        tipo = sheet.cell(row=r, column=2).value
        if tipo not in ('variacao', 'sem-variacao'):
            continue

        total_linhas += 1
        sku = str(sheet.cell(row=r, column=4).value or '').strip()
        if not sku:
            continue

        sku_pai = str(sheet.cell(row=r, column=3).value or '').strip() if tipo == 'variacao' else ''
        parent = parents.get(sku_pai, {}) if tipo == 'variacao' else {}

        nome_original = str(sheet.cell(row=r, column=11).value or '').strip()
        nome_pai = parent.get('nome', '') if parent else ''
        nome = nome_original if nome_original else (nome_pai if nome_pai else sku)

        img = str(sheet.cell(row=r, column=37).value or '').strip()
        if not img and parent:
            img = parent.get('imagem_url', '')

        cat1 = str(sheet.cell(row=r, column=32).value or '').strip()
        cat2 = str(sheet.cell(row=r, column=33).value or '').strip()
        if not cat1 and parent:
            cat1 = parent.get('cat1', '')
            cat2 = parent.get('cat2', '')

        categoria = (cat1 + (' > ' + cat2 if cat2 else '')).strip()
        if not categoria:
            categoria = 'Sem Categoria'

        # Grades (tamanho, cor, etc.)
        tam = str(sheet.cell(row=r, column=54).value or '').strip()
        cor = str(sheet.cell(row=r, column=44).value or '').strip()
        acess = str(sheet.cell(row=r, column=53).value or '').strip()
        var_nome = [g for g in [tam, cor, acess] if g]
        variacao = ' / '.join(var_nome) if var_nome else ('Padrão' if tipo == 'sem-variacao' else 'Único')

        # Estoque (Coluna 17)
        try:
            estoque = int(float(sheet.cell(row=r, column=17).value or 0))
        except (ValueError, TypeError):
            estoque = 0

        # Custo (Coluna 21)
        try:
            custo = float(sheet.cell(row=r, column=21).value or 0.0)
        except (ValueError, TypeError):
            custo = 0.0

        # Preço Cheio (Coluna 23)
        try:
            preco_cheio = float(sheet.cell(row=r, column=23).value or 0.0)
        except (ValueError, TypeError):
            preco_cheio = 0.0

        # Preço Promocional (Coluna 24)
        try:
            preco_promo = float(sheet.cell(row=r, column=24).value or 0.0)
        except (ValueError, TypeError):
            preco_promo = 0.0

        preco_venda = preco_promo if preco_promo > 0 else preco_cheio

        ativo = str(sheet.cell(row=r, column=5).value or 'S')
        ncm = str(sheet.cell(row=r, column=8).value or '').strip()

        try:
            peso_g = float(sheet.cell(row=r, column=20).value or 0.0)
        except (ValueError, TypeError):
            peso_g = 0.0

        item = {
            'sku': sku,
            'sku_pai': sku_pai if sku_pai else sku,
            'nome': nome,
            'variacao': variacao,
            'tamanho': tam,
            'cor': cor,
            'categoria': categoria,
            'custo': custo,
            'preco_cheio': preco_cheio,
            'preco_promo': preco_promo,
            'preco_venda': preco_venda,
            'estoque': estoque,
            'ativo': ativo,
            'imagem_url': img,
            'peso_g': peso_g,
            'ncm': ncm
        }
        items.append(item)

        # Se for item sem-variacao que nao estava em parents, adiciona como pai solo
        if not sku_pai and sku not in parents:
            parents[sku] = {
                'sku_pai': sku,
                'nome': nome,
                'ativo': ativo,
                'destaque': 'N',
                'ncm': ncm,
                'categoria': categoria,
                'cat1': cat1,
                'cat2': cat2,
                'imagem_url': img,
                'imagem_url2': '',
                'descricao': '',
                'variacoes': []
            }

    # 3. Consolidação de agregados por produto pai
    items_by_parent = {}
    for it in items:
        p_sku = it['sku_pai']
        if p_sku not in items_by_parent:
            items_by_parent[p_sku] = []
        items_by_parent[p_sku].append(it)

    # Ingestão com Upsert nas tabelas produtos_pai e produtos_variacao
    sucessos_pai = 0
    sucessos_var = 0

    for p_sku, p_info in parents.items():
        vars_list = items_by_parent.get(p_sku, [])
        if vars_list:
            custos = [v['custo'] for v in vars_list if v['custo'] > 0] or [0.0]
            precos = [v['preco_venda'] for v in vars_list if v['preco_venda'] > 0] or [0.0]
            custo_min = min(custos)
            custo_max = max(custos)
            preco_venda_min = min(precos)
            preco_venda_max = max(precos)
            estoque_total = sum(v['estoque'] for v in vars_list)
        else:
            custo_min = custo_max = preco_venda_min = preco_venda_max = 0.0
            estoque_total = 0

        cursor.execute("""
        INSERT INTO produtos_pai (
            sku_pai, nome, categoria, cat1, cat2, imagem_url, imagem_url2, ncm, descricao,
            custo_min, custo_max, preco_venda_min, preco_venda_max, estoque_total, ativo, destaque, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(sku_pai) DO UPDATE SET
            nome = excluded.nome,
            categoria = excluded.categoria,
            cat1 = excluded.cat1,
            cat2 = excluded.cat2,
            imagem_url = excluded.imagem_url,
            imagem_url2 = excluded.imagem_url2,
            ncm = excluded.ncm,
            descricao = excluded.descricao,
            custo_min = excluded.custo_min,
            custo_max = excluded.custo_max,
            preco_venda_min = excluded.preco_venda_min,
            preco_venda_max = excluded.preco_venda_max,
            estoque_total = excluded.estoque_total,
            ativo = excluded.ativo,
            destaque = excluded.destaque,
            updated_at = CURRENT_TIMESTAMP
        """, (
            p_sku, p_info['nome'], p_info['categoria'], p_info['cat1'], p_info['cat2'],
            p_info['imagem_url'], p_info['imagem_url2'], p_info['ncm'], p_info['descricao'],
            custo_min, custo_max, preco_venda_min, preco_venda_max, estoque_total, p_info['ativo'], p_info['destaque']
        ))
        sucessos_pai += 1

    # Obter IDs dos pais inseridos
    cursor.execute("SELECT id, sku_pai FROM produtos_pai")
    pai_id_map = {row['sku_pai']: row['id'] for row in cursor.fetchall()}

    for it in items:
        pai_id = pai_id_map.get(it['sku_pai'])
        cursor.execute("""
        INSERT INTO produtos_variacao (
            produto_pai_id, sku, sku_pai, nome, variacao, categoria, tamanho, cor,
            custo, preco_cheio, preco_promo, preco_venda, estoque, ativo, imagem_url, peso_g, ncm, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(sku) DO UPDATE SET
            produto_pai_id = excluded.produto_pai_id,
            sku_pai = excluded.sku_pai,
            nome = excluded.nome,
            variacao = excluded.variacao,
            categoria = excluded.categoria,
            tamanho = excluded.tamanho,
            cor = excluded.cor,
            custo = excluded.custo,
            preco_cheio = excluded.preco_cheio,
            preco_promo = excluded.preco_promo,
            preco_venda = excluded.preco_venda,
            estoque = excluded.estoque,
            ativo = excluded.ativo,
            imagem_url = excluded.imagem_url,
            peso_g = excluded.peso_g,
            ncm = excluded.ncm,
            updated_at = CURRENT_TIMESTAMP
        """, (
            pai_id, it['sku'], it['sku_pai'], it['nome'], it['variacao'], it['categoria'],
            it['tamanho'], it['cor'], it['custo'], it['preco_cheio'], it['preco_promo'],
            it['preco_venda'], it['estoque'], it['ativo'], it['imagem_url'], it['peso_g'], it['ncm']
        ))
        sucessos_var += 1

    # Registrar log de importação
    cursor.execute("""
    INSERT INTO historico_importacoes (tipo_origem, nome_arquivo, total_registros, registros_sucesso, registros_erros, detalhes_json)
    VALUES (?, ?, ?, ?, ?, ?)
    """, (
        'CATALOGO_LOJA',
        nome_arquivo,
        total_linhas,
        sucessos_var,
        0,
        json.dumps({'produtos_pai': sucessos_pai, 'produtos_variacao': sucessos_var}, ensure_ascii=False)
    ))

    conn.commit()
    conn.close()

    print(f"✅ Catálogo importado com sucesso: {sucessos_pai} modelos pai e {sucessos_var} variações no SQLite.")
    return True


if __name__ == '__main__':
    importar_catalogo()
