# -*- coding: utf-8 -*-
"""
Módulo para ingestão e auditoria dos anúncios da Shopee no SQLite.
Cruza anúncios exportados com o catálogo de produtos e salva métricas 2026.
"""

import os
import io
import sys
import glob
import re
import json
import difflib
import unicodedata
import zipfile
import openpyxl

from .database import get_connection, init_db, ROOT

PASTA_SHOPEE = os.path.join(ROOT, 'data', 'raw', 'shopee')
MAP_HELPER_PATH = os.path.join(ROOT, 'scripts', 'config', 'map_helper.txt')
EXACT_MAP_PATH = os.path.join(ROOT, 'scripts', 'utils', 'shopee_exact_map.py')

EMBALAGEM_PADRAO = 1.50


def carregar_workbook_seguro(caminho):
    """Carrega o workbook contornando eventuais erros de tag XML de pane da exportação Shopee."""
    with zipfile.ZipFile(caminho, 'r') as z_in:
        out_buf = io.BytesIO()
        with zipfile.ZipFile(out_buf, 'w') as z_out:
            for item in z_in.infolist():
                data = z_in.read(item.filename)
                if item.filename.endswith('sheet1.xml'):
                    data = data.replace(b'activePane="bottom_left"', b'activePane="bottomLeft"')
                z_out.writestr(item, data)
    out_buf.seek(0)
    return openpyxl.load_workbook(out_buf, data_only=True)


def normalizar_texto(txt):
    if not txt:
        return ''
    txt = unicodedata.normalize('NFKD', str(txt)).encode('ASCII', 'ignore').decode('ASCII')
    txt = re.sub(r'[^\w\s]', ' ', txt.lower())
    return ' '.join(txt.split())


def id_seguro(valor):
    if valor is None:
        return ''
    try:
        numerico = float(valor)
    except (TypeError, ValueError):
        return str(valor).strip()
    if not numerico.is_integer():
        return str(valor).strip()
    return str(int(numerico))


def calc_shopee_fee(p):
    """Calcula taxas Shopee 2026 (Artigo 26839)."""
    if p <= 0:
        return 0.0, 'Inválido'
    if p < 9.00:
        return (p * 0.20 + p * 0.50), 'Sub R$ 9: 20% + 50% taxa fixa'
    if p <= 79.99:
        return (p * 0.20 + 4.50), 'Até R$ 79,99: 20% + R$ 4,50'
    if p <= 99.99:
        return (p * 0.14 + 16.00), 'R$ 80 a 99,99: 14% + R$ 16,00'
    if p <= 199.99:
        return (p * 0.14 + 20.00), 'R$ 100 a 199,99: 14% + R$ 20,00'
    return (p * 0.14 + 26.00), 'Acima de R$ 200: 14% + R$ 26,00'


def calc_max_safe_discount(cad_price, cost, pack, target_margin):
    if cad_price <= 0:
        return 0
    m_target = target_margin / 100.0
    for disc in range(90, -1, -1):
        p_sim = cad_price * (1.0 - (disc / 100.0))
        fee, _ = calc_shopee_fee(p_sim)
        profit = p_sim - fee - cost - pack
        margin = (profit / p_sim) if p_sim > 0 else 0
        if margin >= m_target and profit >= 0:
            return disc
    return 0


def carregar_map_helper():
    """Carrega o arquivo de mapeamento manual se existir."""
    id_map = {}
    if os.path.exists(MAP_HELPER_PATH):
        with open(MAP_HELPER_PATH, 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith('#') and '->' in line:
                    parts = line.split('->')
                    if len(parts) >= 2:
                        k = parts[0].strip()
                        v = parts[1].strip()
                        id_map[k] = v
    return id_map


def carregar_exact_map():
    """Carrega o dicionário EXACT_PROD_MAP de utils/shopee_exact_map.py se existir."""
    if os.path.exists(EXACT_MAP_PATH):
        try:
            sys.path.insert(0, os.path.dirname(EXACT_MAP_PATH))
            from shopee_exact_map import EXACT_PROD_MAP
            return EXACT_PROD_MAP
        except Exception:
            pass
    return {}


def importar_shopee(caminho_arquivo=None, db_path=None):
    """Importa os anúncios da Shopee, realiza o matching com produtos_variacao e grava no SQLite."""
    init_db(db_path)
    conn = get_connection(db_path)
    cursor = conn.cursor()

    if not caminho_arquivo:
        padroes = ['Shopee_Anuncios_Exportacao*.xlsx', 'mass_update_sales_info*.xlsx']
        for padrao in padroes:
            encontrados = glob.glob(os.path.join(PASTA_SHOPEE, padrao))
            if encontrados:
                encontrados.sort(key=os.path.getmtime, reverse=True)
                caminho_arquivo = encontrados[0]
                break

    if not caminho_arquivo or not os.path.exists(caminho_arquivo):
        print(f"Erro: Nenhuma planilha Shopee encontrada em {os.path.relpath(PASTA_SHOPEE, ROOT)}.")
        conn.close()
        return False

    nome_arquivo = os.path.basename(caminho_arquivo)
    print(f"🛍️ Importando anúncios Shopee: {nome_arquivo}...")

    # Carregar produtos do banco para matching
    cursor.execute("""
    SELECT v.sku, v.sku_pai, v.nome, v.variacao, v.categoria, v.custo, v.preco_venda, v.estoque, v.imagem_url
    FROM produtos_variacao v
    """)
    variacoes_db = [dict(row) for row in cursor.fetchall()]

    sku_map = {v['sku'].lower(): v for v in variacoes_db if v['sku']}
    parent_sku_map = {}
    for v in variacoes_db:
        p_sku = (v.get('sku_pai') or '').lower()
        if p_sku:
            if p_sku not in parent_sku_map:
                parent_sku_map[p_sku] = []
            parent_sku_map[p_sku].append(v)

    map_helper = carregar_map_helper()
    exact_map = carregar_exact_map()

    # Leitura da planilha Shopee
    wb = carregar_workbook_seguro(caminho_arquivo)
    sheet = wb.active

    is_anuncios = nome_arquivo.lower().startswith('shopee_anuncios_exportacao')
    start_row = 2 if is_anuncios else 7

    shopee_items = []
    total_linhas = 0

    for r in range(start_row, sheet.max_row + 1):
        if is_anuncios:
            p_id = id_seguro(sheet.cell(row=r, column=1).value)
            p_name = str(sheet.cell(row=r, column=2).value or '').strip()
            p_sku = str(sheet.cell(row=r, column=3).value or '').strip()
            v_name1 = str(sheet.cell(row=r, column=4).value or '').strip()
            v_opt1 = str(sheet.cell(row=r, column=5).value or '').strip()
            v_sku = str(sheet.cell(row=r, column=8).value or '').strip()
            v_id = id_seguro(sheet.cell(row=r, column=9).value)

            try:
                cad_price = float(sheet.cell(row=r, column=14).value or 0.0)
            except (ValueError, TypeError):
                cad_price = 0.0

            try:
                promo_price = float(sheet.cell(row=r, column=15).value or 0.0)
            except (ValueError, TypeError):
                promo_price = 0.0

            try:
                stock = int(float(sheet.cell(row=r, column=16).value or 0))
            except (ValueError, TypeError):
                stock = 0

            img = str(sheet.cell(row=r, column=17).value or '').strip()
            var_full = v_opt1 if v_opt1 else v_name1

        else:
            # Layout legado (mass_update_sales_info)
            p_id = id_seguro(sheet.cell(row=r, column=1).value)
            p_name = str(sheet.cell(row=r, column=2).value or '').strip()
            p_sku = str(sheet.cell(row=r, column=3).value or '').strip()
            var_full = str(sheet.cell(row=r, column=4).value or '').strip()
            v_sku = str(sheet.cell(row=r, column=5).value or '').strip()
            v_id = id_seguro(sheet.cell(row=r, column=6).value)

            try:
                cad_price = float(sheet.cell(row=r, column=7).value or 0.0)
            except (ValueError, TypeError):
                cad_price = 0.0

            promo_price = cad_price
            stock = 0
            img = ''

        if not p_name and not v_sku and not p_sku and not p_id:
            continue

        total_linhas += 1

        if promo_price <= 0:
            promo_price = cad_price
        if cad_price <= 0:
            cad_price = promo_price

        real_disc_pct = ((cad_price - promo_price) / cad_price * 100.0) if cad_price > 0 else 0.0
        if real_disc_pct < 0:
            real_disc_pct = 0.0

        # Algoritmo de Matching
        matched_item = None
        match_type = 'Não Vinculado'

        # 1. Matching por SKU Direto
        if v_sku and v_sku.lower() in sku_map:
            matched_item = sku_map[v_sku.lower()]
            match_type = 'SKU Direto'
        elif p_sku and p_sku.lower() in sku_map:
            matched_item = sku_map[p_sku.lower()]
            match_type = 'SKU Principal'

        # 2. Matching por Dicionário Manual
        if not matched_item:
            target_sku = map_helper.get(v_id) or map_helper.get(v_sku) or exact_map.get(v_id) or exact_map.get(v_sku)
            if target_sku and target_sku.lower() in sku_map:
                matched_item = sku_map[target_sku.lower()]
                match_type = 'ID Variante (Mapeado)'

        # 3. Matching por SKU Pai + Variação
        if not matched_item and p_sku and p_sku.lower() in parent_sku_map:
            candidates = parent_sku_map[p_sku.lower()]
            if len(candidates) == 1:
                matched_item = candidates[0]
                match_type = 'SKU Pai (Único)'
            else:
                v_norm = normalizar_texto(var_full)
                for c in candidates:
                    if normalizar_texto(c['variacao']) in v_norm or v_norm in normalizar_texto(c['variacao']):
                        matched_item = c
                        match_type = 'SKU Pai + Variação'
                        break
                if not matched_item:
                    matched_item = candidates[0]
                    match_type = 'SKU Pai (Padrão)'

        # 4. Matching Fuzzy por Similaridade de Título
        if not matched_item:
            p_norm = normalizar_texto(p_name)
            melhor_score = 0.0
            melhor_candidato = None
            for v in variacoes_db:
                v_title_norm = normalizar_texto(v['nome'])
                ratio = difflib.SequenceMatcher(None, p_norm, v_title_norm).ratio()
                if ratio > melhor_score:
                    melhor_score = ratio
                    melhor_candidato = v

            if melhor_score >= 0.60:
                matched_item = melhor_candidato
                match_type = f'Similaridade ({int(melhor_score*100)}%)'

        # Atribuição dos dados vinculados
        custo = matched_item['custo'] if matched_item else 0.0
        preco_loja = matched_item['preco_venda'] if matched_item else 0.0
        loja_sku = matched_item['sku'] if matched_item else None
        if not img and matched_item:
            img = matched_item.get('imagem_url', '')

        # Cálculos de Auditoria Shopee 2026
        fee, tier = calc_shopee_fee(promo_price)
        repasse = promo_price - fee
        profit_real = promo_price - fee - custo - EMBALAGEM_PADRAO
        margin_real = (profit_real / promo_price * 100.0) if promo_price > 0 else 0.0

        desc_20 = calc_max_safe_discount(cad_price, custo, EMBALAGEM_PADRAO, 20.0)
        desc_10 = calc_max_safe_discount(cad_price, custo, EMBALAGEM_PADRAO, 10.0)

        shopee_items.append({
            'shopee_item_id': p_id,
            'shopee_var_id': v_id,
            'nome_anuncio': p_name,
            'nome_variacao': var_full,
            'sku_anuncio': p_sku,
            'sku_variacao': v_sku,
            'sku_produto_loja': loja_sku,
            'match_tipo': match_type,
            'preco_cadastro_ancora': round(cad_price, 2),
            'preco_real_promo': round(promo_price, 2),
            'desconto_real_pct': round(real_disc_pct, 1),
            'estoque_shopee': stock,
            'imagem_url': img,
            'taxa_estimada': round(fee, 2),
            'repasse_estimado': round(repasse, 2),
            'lucro_estimado': round(profit_real, 2),
            'margem_estimada': round(margin_real, 1),
            'tier_shopee': tier,
            'desc_max_seguro_20': desc_20,
            'desc_max_seguro_10': desc_10
        })

    # Limpar e recarregar tabela anuncios_shopee
    cursor.execute("DELETE FROM anuncios_shopee;")

    cursor.executemany("""
    INSERT INTO anuncios_shopee (
        shopee_item_id, shopee_var_id, nome_anuncio, nome_variacao, sku_anuncio, sku_variacao,
        sku_produto_loja, match_tipo, preco_cadastro_ancora, preco_real_promo, desconto_real_pct,
        estoque_shopee, imagem_url, taxa_estimada, repasse_estimado, lucro_estimado, margem_estimada,
        tier_shopee, desc_max_seguro_20, desc_max_seguro_10, ultima_sincronizacao
    ) VALUES (
        :shopee_item_id, :shopee_var_id, :nome_anuncio, :nome_variacao, :sku_anuncio, :sku_variacao,
        :sku_produto_loja, :match_tipo, :preco_cadastro_ancora, :preco_real_promo, :desconto_real_pct,
        :estoque_shopee, :imagem_url, :taxa_estimada, :repasse_estimado, :lucro_estimado, :margem_estimada,
        :tier_shopee, :desc_max_seguro_20, :desc_max_seguro_10, CURRENT_TIMESTAMP
    )
    """, shopee_items)

    # Log de importação
    cursor.execute("""
    INSERT INTO historico_importacoes (tipo_origem, nome_arquivo, total_registros, registros_sucesso, registros_erros, detalhes_json)
    VALUES (?, ?, ?, ?, ?, ?)
    """, (
        'SHOPEE_ANUNCIOS',
        nome_arquivo,
        total_linhas,
        len(shopee_items),
        0,
        json.dumps({'itens_importados': len(shopee_items)}, ensure_ascii=False)
    ))

    conn.commit()
    conn.close()

    print(f"✅ Anúncios Shopee importados com sucesso: {len(shopee_items)} itens gravados no SQLite.")
    return True


if __name__ == '__main__':
    importar_shopee()
