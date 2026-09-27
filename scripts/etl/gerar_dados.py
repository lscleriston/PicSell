# -*- coding: utf-8 -*-
"""
Script para processar a planilha de produtos exportada da Loja Integrada
e gerar os arquivos de dados para a aplicação web (HTML/JS).
"""

import os
import sys
import glob
import json
import re
import html
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PASTA_CATALOGO = os.path.join(ROOT, 'data', 'raw', 'catalogo')
SAIDA_JSON = os.path.join(ROOT, 'data', 'processed', 'dados_produtos.json')
SAIDA_JS = os.path.join(ROOT, 'web', 'data', 'dados_produtos.js')

def clean_html(raw_html):
    if not raw_html:
        return ''
    cleanr = re.compile(r'<.*?>')
    cleantext = re.sub(cleanr, ' ', str(raw_html))
    return ' '.join(html.unescape(cleantext).split()).strip()

def processar_planilha(caminho_arquivo=None):
    if not caminho_arquivo:
        # Busca o arquivo xlsx mais recente na pasta de catalogo
        arquivos = glob.glob(os.path.join(PASTA_CATALOGO, '*.xlsx'))
        if not arquivos:
            print(f"Erro: Nenhum arquivo .xlsx encontrado em {os.path.relpath(PASTA_CATALOGO, ROOT)}.")
            return None
        # Ordena pelo mais recente
        arquivos.sort(key=os.path.getmtime, reverse=True)
        caminho_arquivo = arquivos[0]

    print(f"Lendo planilha: {caminho_arquivo}...")
    wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
    sheet = wb.active

    # Mapeamento dos produtos pai (com-variacao)
    parents = {}
    for r in range(3, sheet.max_row + 1):
        tipo = sheet.cell(row=r, column=2).value
        if tipo == 'com-variacao':
            sku = str(sheet.cell(row=r, column=4).value or '').strip()
            cat1 = str(sheet.cell(row=r, column=32).value or '').strip()
            cat2 = str(sheet.cell(row=r, column=33).value or '').strip()
            categoria = cat1 + (' > ' + cat2 if cat2 else '')
            parents[sku] = {
                'id': sheet.cell(row=r, column=1).value,
                'sku': sku,
                'nome': str(sheet.cell(row=r, column=11).value or '').strip(),
                'ativo': sheet.cell(row=r, column=5).value or 'S',
                'destaque': sheet.cell(row=r, column=7).value or 'N',
                'ncm': str(sheet.cell(row=r, column=8).value or '').strip(),
                'categoria': categoria,
                'cat1': cat1,
                'cat2': cat2,
                'img': str(sheet.cell(row=r, column=37).value or '').strip(),
                'img2': str(sheet.cell(row=r, column=38).value or '').strip(),
                'descricao': clean_html(sheet.cell(row=r, column=14).value)[:300],
                'variacoes': []
            }

    # Processamento dos itens (variacao e sem-variacao)
    items = []
    categorias_set = set()

    for r in range(3, sheet.max_row + 1):
        tipo = sheet.cell(row=r, column=2).value
        if tipo not in ('variacao', 'sem-variacao'):
            continue

        sku = str(sheet.cell(row=r, column=4).value or '').strip()
        sku_pai = str(sheet.cell(row=r, column=3).value or '').strip() if tipo == 'variacao' else ''
        parent = parents.get(sku_pai, {}) if tipo == 'variacao' else {}

        nome_original = str(sheet.cell(row=r, column=11).value or '').strip()
        nome_pai = parent.get('nome', '') if parent else ''
        nome = nome_original if nome_original else (nome_pai if nome_pai else sku)

        img = str(sheet.cell(row=r, column=37).value or '').strip()
        if not img and parent:
            img = parent.get('img', '')

        cat1 = str(sheet.cell(row=r, column=32).value or '').strip()
        cat2 = str(sheet.cell(row=r, column=33).value or '').strip()
        if not cat1 and parent:
            cat1 = parent.get('cat1', '')
            cat2 = parent.get('cat2', '')

        categoria = (cat1 + (' > ' + cat2 if cat2 else '')).strip()
        if not categoria:
            categoria = 'Sem Categoria'
        categorias_set.add(categoria)

        # Grades (tamanho, cor, etc.)
        tam = str(sheet.cell(row=r, column=54).value or '').strip()
        cor = str(sheet.cell(row=r, column=44).value or '').strip()
        acess = str(sheet.cell(row=r, column=53).value or '').strip()
        grades = [g for g in [tam, cor, acess] if g]
        variacao_nome = ' / '.join(grades) if grades else ('Padrão' if tipo == 'sem-variacao' else 'Único')

        # Estoque
        stock = sheet.cell(row=r, column=17).value
        try:
            stock = float(stock) if stock is not None else 0.0
        except:
            stock = 0.0

        # Custo
        cost = sheet.cell(row=r, column=21).value
        try:
            cost = float(cost) if cost is not None else 0.0
        except:
            cost = 0.0

        # Preço Cheio
        price = sheet.cell(row=r, column=23).value
        try:
            price = float(price) if price is not None else 0.0
        except:
            price = 0.0

        # Preço Promocional
        promo = sheet.cell(row=r, column=24).value
        try:
            promo = float(promo) if promo is not None else 0.0
        except:
            promo = 0.0

        preco_venda = promo if promo > 0 else price
        custo_total = stock * cost
        venda_total = stock * preco_venda
        lucro_total = venda_total - custo_total
        margem_pct = (lucro_total / venda_total * 100.0) if venda_total > 0 else (
            ((preco_venda - cost) / preco_venda * 100.0) if preco_venda > 0 else 0.0
        )

        ativo = sheet.cell(row=r, column=5).value or 'S'

        item = {
            'id': sheet.cell(row=r, column=1).value,
            'sku': sku,
            'sku_pai': sku_pai,
            'tipo': tipo,
            'nome': nome,
            'nome_pai': nome_pai,
            'variacao': variacao_nome,
            'categoria': categoria,
            'img': img,
            'ativo': ativo,
            'estoque': round(stock, 2),
            'custo': round(cost, 2),
            'preco_cheio': round(price, 2),
            'preco_promo': round(promo, 2),
            'preco_venda': round(preco_venda, 2),
            'custo_total': round(custo_total, 2),
            'venda_total': round(venda_total, 2),
            'lucro_total': round(lucro_total, 2),
            'margem_pct': round(margem_pct, 1)
        }
        items.append(item)

        if sku_pai and sku_pai in parents:
            parents[sku_pai]['variacoes'].append(item)

    # Agrupamento por produto (para a visão consolidada)
    produtos_agrupados = []

    # Primeiro, adiciona os pais com suas variações
    for p_sku, p in parents.items():
        vars_list = p['variacoes']
        if not vars_list:
            continue

        estoque_total = sum(v['estoque'] for v in vars_list)
        custo_total = sum(v['custo_total'] for v in vars_list)
        venda_total = sum(v['venda_total'] for v in vars_list)
        lucro_total = venda_total - custo_total
        margem_pct = (lucro_total / venda_total * 100.0) if venda_total > 0 else 0.0

        custos = [v['custo'] for v in vars_list if v['custo'] > 0]
        precos = [v['preco_venda'] for v in vars_list if v['preco_venda'] > 0]

        custo_min = min(custos) if custos else 0.0
        custo_max = max(custos) if custos else 0.0
        preco_min = min(precos) if precos else 0.0
        preco_max = max(precos) if precos else 0.0

        produtos_agrupados.append({
            'sku': p_sku,
            'tipo': 'com-variacao',
            'nome': p['nome'],
            'categoria': p['categoria'] or 'Sem Categoria',
            'img': p['img'],
            'ativo': p['ativo'],
            'descricao': p['descricao'],
            'estoque_total': round(estoque_total, 2),
            'custo_total': round(custo_total, 2),
            'venda_total': round(venda_total, 2),
            'lucro_total': round(lucro_total, 2),
            'margem_pct': round(margem_pct, 1),
            'custo_min': round(custo_min, 2),
            'custo_max': round(custo_max, 2),
            'preco_min': round(preco_min, 2),
            'preco_max': round(preco_max, 2),
            'qtd_variacoes': len(vars_list),
            'com_estoque': any(v['estoque'] > 0 for v in vars_list),
            'variacoes': vars_list
        })

    # Adiciona os sem-variacao como produtos agrupados individuais
    for it in items:
        if it['tipo'] == 'sem-variacao':
            produtos_agrupados.append({
                'sku': it['sku'],
                'tipo': 'sem-variacao',
                'nome': it['nome'],
                'categoria': it['categoria'],
                'img': it['img'],
                'ativo': it['ativo'],
                'descricao': '',
                'estoque_total': it['estoque'],
                'custo_total': it['custo_total'],
                'venda_total': it['venda_total'],
                'lucro_total': it['lucro_total'],
                'margem_pct': it['margem_pct'],
                'custo_min': it['custo'],
                'custo_max': it['custo'],
                'preco_min': it['preco_venda'],
                'preco_max': it['preco_venda'],
                'qtd_variacoes': 1,
                'com_estoque': it['estoque'] > 0,
                'variacoes': [it]
            })

    # Resumo Geral / KPIs
    total_estoque = sum(x['estoque'] for x in items)
    total_custo = sum(x['custo_total'] for x in items)
    total_venda = sum(x['venda_total'] for x in items)
    total_lucro = total_venda - total_custo
    margem_geral = (total_lucro / total_venda * 100.0) if total_venda > 0 else 0.0

    kpis = {
        'total_skus': len(items),
        'total_produtos': len(produtos_agrupados),
        'total_pecas_estoque': round(total_estoque, 0),
        'skus_com_estoque': sum(1 for x in items if x['estoque'] > 0),
        'skus_sem_estoque': sum(1 for x in items if x['estoque'] == 0),
        'skus_estoque_baixo': sum(1 for x in items if 0 < x['estoque'] <= 2),
        'total_custo_estoque': round(total_custo, 2),
        'total_venda_estoque': round(total_venda, 2),
        'lucro_bruto_estoque': round(total_lucro, 2),
        'margem_media_estoque': round(margem_geral, 1),
        'arquivo_origem': os.path.basename(caminho_arquivo)
    }

    resultado = {
        'kpis': kpis,
        'categorias': sorted(list(categorias_set)),
        'produtos_agrupados': produtos_agrupados,
        'itens_detalhados': items
    }

    # Salva em JSON
    os.makedirs(os.path.dirname(SAIDA_JSON), exist_ok=True)
    os.makedirs(os.path.dirname(SAIDA_JS), exist_ok=True)
    with open(SAIDA_JSON, 'w', encoding='utf-8') as f:
        json.dump(resultado, f, ensure_ascii=False, indent=2)

    # Salva em JS com window.DADOS_PRODUTOS para compatibilidade direta no navegador
    with open(SAIDA_JS, 'w', encoding='utf-8') as f:
        f.write('window.DADOS_PRODUTOS = ' + json.dumps(resultado, ensure_ascii=False) + ';\n')

    print("\nProcessamento concluído com sucesso!")
    print(f"- Total de SKUs: {kpis['total_skus']}")
    print(f"- Produtos Consolidados: {kpis['total_produtos']}")
    print(f"- Peças em Estoque: {kpis['total_pecas_estoque']}")
    print(f"- Custo Total em Estoque: R$ {kpis['total_custo_estoque']:,.2f}")
    print(f"- Venda Total Estimada: R$ {kpis['total_venda_estoque']:,.2f}")
    print(f"- Lucro Bruto Estimado: R$ {kpis['lucro_bruto_estoque']:,.2f} ({kpis['margem_media_estoque']}%)")
    print(f"- Arquivos gerados: {os.path.relpath(SAIDA_JSON, ROOT)} e {os.path.relpath(SAIDA_JS, ROOT)}")

    return resultado

if __name__ == '__main__':
    caminho = sys.argv[1] if len(sys.argv) > 1 else None
    processar_planilha(caminho)
