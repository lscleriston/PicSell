# -*- coding: utf-8 -*-
"""
ETL da auditoria Shopee (consolidacao dos antigos analisar_shopee.py e scratch_match.py).

Entrada : data/raw/shopee/*.xlsx + data/processed/dados_produtos.json
Saida   : web/data/dados_shopee_real.js + data/processed/analise_shopee_real.json

Formatos aceitos em data/raw/shopee/:
  - Shopee_Anuncios_Exportacao*.xlsx  (exportacao de anuncios, preco cadastro + promocao)
  - mass_update_sales_info*.xlsx      (exportacao legada de precos, inicia na linha 7)
"""

import os
import io
import sys
import re
import json
import glob
import difflib
import unicodedata
import zipfile

import openpyxl

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(ROOT, 'scripts', 'utils'))

from shopee_exact_map import EXACT_PROD_MAP

RAW_DIR = os.path.join(ROOT, 'data', 'raw', 'shopee')
PROC_DIR = os.path.join(ROOT, 'data', 'processed')
WEB_DATA_DIR = os.path.join(ROOT, 'web', 'data')

ARQUIVO_PRODUTOS_JSON = os.path.join(PROC_DIR, 'dados_produtos.json')
ARQUIVO_PRODUTOS_JS = os.path.join(WEB_DATA_DIR, 'dados_produtos.js')
ARQUIVO_AUDITORIA_JS = os.path.join(WEB_DATA_DIR, 'dados_shopee_real.js')
ARQUIVO_AUDITORIA_JSON = os.path.join(PROC_DIR, 'analise_shopee_real.json')

EMBALAGEM = 1.50


def relativo(caminho):
    return os.path.relpath(caminho, ROOT)


def carregar_workbook(caminho):
    """Carrega a planilha contornando o problema de activePane do arquivo exportado."""
    with zipfile.ZipFile(caminho, 'r') as z_in:
        out_buf = io.BytesIO()
        with zipfile.ZipFile(out_buf, 'w') as z_out:
            for item in z_in.infolist():
                data = z_in.read(item.filename)
                if item.filename.endswith('sheet1.xml'):
                    data = data.replace(b'activePane="bottom_left"', b'activePane="bottomLeft"')
                z_out.writestr(item, data)
    out_buf.seek(0)
    return openpyxl.load_workbook(out_buf, data_only=True)


def localizar_planilha():
    """Prioriza a exportacao atual de anuncios; usa o mass_update como historico."""
    padroes = ['Shopee_Anuncios_Exportacao*.xlsx', 'mass_update_sales_info*.xlsx']
    for padrao in padroes:
        encontrados = glob.glob(os.path.join(RAW_DIR, padrao))
        if encontrados:
            encontrados.sort(key=os.path.getmtime, reverse=True)
            return encontrados[0]
    return None


def eh_exportacao_anuncios(caminho):
    return os.path.basename(caminho).lower().startswith('shopee_anuncios_exportacao')


def para_float(valor):
    try:
        return float(valor)
    except (TypeError, ValueError):
        return 0.0



def id_seguro(valor):
    """Normaliza IDs (que chegam como float na exportacao) para texto sem casas decimais."""
    if valor is None:
        return ''
    try:
        numerico = float(valor)
    except (TypeError, ValueError):
        return str(valor).strip()
    if not numerico.is_integer():
        return str(valor).strip()
    return str(int(numerico))


def ler_exportacao_anuncios(sheet):
    linhas = list(sheet.iter_rows(values_only=True))
    dados = [r for r in linhas[1:] if any(r)]
    itens = []
    for r in dados:
        var_nome_1 = str(r[7] or '').strip() if r[7] is not None else ''
        var_nome_2 = str(r[10] or '').strip() if r[10] is not None else ''
        var_name = f"{var_nome_1} {var_nome_2}".strip() if (var_nome_1 and var_nome_2) else (var_nome_1 or var_nome_2)
        preco_cad = para_float(r[13])
        preco_promo = para_float(r[14]) or preco_cad
        itens.append({
            'prod_id': id_seguro(r[1]),
            'prod_name': str(r[2] or '').strip(),
            'parent_sku': str(r[3] or '').strip(),
            'var_name': var_name,
            'var_sku': str(r[11] or '').strip(),
            'var_id': id_seguro(r[12]),
            'preco_cad': preco_cad,
            'preco_promo': preco_promo,
            'qtd': para_float(r[15]),
            'img': str(r[8] or '').strip() or str(r[19] or '').strip(),
        })
    return itens


def ler_mass_update(sheet):
    itens = []
    for linha in range(7, sheet.max_row + 1):
        prod_id = sheet.cell(row=linha, column=1).value
        preco = para_float(sheet.cell(row=linha, column=7).value)
        if not prod_id and not preco:
            continue
        itens.append({
            'prod_id': id_seguro(prod_id),
            'prod_name': str(sheet.cell(row=linha, column=2).value or '').strip(),
            'parent_sku': str(sheet.cell(row=linha, column=5).value or '').strip(),
            'var_name': str(sheet.cell(row=linha, column=4).value or '').strip(),
            'var_sku': str(sheet.cell(row=linha, column=6).value or '').strip(),
            'var_id': id_seguro(sheet.cell(row=linha, column=3).value),
            'preco_cad': preco,
            'preco_promo': preco,
            'qtd': para_float(sheet.cell(row=linha, column=9).value),
            'img': '',
        })
    return itens


def carregar_json(caminho):
    if not os.path.exists(caminho):
        return None
    with open(caminho, 'r', encoding='utf-8') as f:
        return json.load(f)


def carregar_janela(caminho, variavel):
    """Le um arquivo .js no formato window.VARIavel = {...};"""
    if not os.path.exists(caminho):
        return None
    with open(caminho, 'r', encoding='utf-8') as f:
        conteudo = f.read()
    busca = re.search(rf'window\.{variavel}\s*=\s*(\{{.*\}});', conteudo, re.DOTALL)
    return json.loads(busca.group(1)) if busca else None


def montar_indice_loja(loja_data):
    itens = loja_data.get('itens_detalhados', [])
    por_sku = {str(it['sku']).strip().upper(): it for it in itens}
    por_nome = {}
    for it in itens:
        por_nome.setdefault(it['nome'].strip().lower(), []).append(it)
    return itens, por_sku, por_nome


def carregar_mapa_anterior():
    dados = carregar_janela(ARQUIVO_AUDITORIA_JS, 'DADOS_SHOPEE_REAL') or {}
    itens = dados.get('itens', [])
    por_var = {}
    por_prod_var = {}
    por_nome = {}
    for it in itens:
        prod_id = str(it.get('prod_id', '')).strip()
        var_id = str(it.get('var_id', '')).strip()
        if var_id:
            por_var[var_id] = it
        if prod_id and var_id:
            por_prod_var[f'{prod_id}_{var_id}'] = it
        if it.get('prod_name'):
            por_nome[por_nome_key(it['prod_name'])] = it
    return por_prod_var, por_var, por_nome


def por_nome_key(texto):
    texto = unicodedata.normalize('NFKD', str(texto)).encode('ASCII', 'ignore').decode('ASCII')
    return texto.lower().strip()


def aplicar_mapeamento_exato(prod_id, var_name, por_nome):
    """Resolve anuncios sem SKU atraves do dicionario de equivalencias manuais."""
    try:
        chave = int(prod_id)
    except (TypeError, ValueError):
        chave = prod_id

    mapping = EXACT_PROD_MAP.get(chave)
    if not mapping:
        return None

    if 'custo_fixo' in mapping:
        return {
            'sku': f'SHOPEE-{chave}',
            'nome': mapping['nome'],
            'variacao': var_name,
            'custo': mapping['custo_fixo'],
            'preco_venda': mapping.get('preco_fixo', 0.0),
            'estoque': 10,
            'img': '',
        }

    candidatos = por_nome.get(mapping['nome'].strip().lower())
    if not candidatos:
        return None

    alvo = candidatos[0]
    for c in candidatos:
        if c.get('variacao') and var_name and (
            var_name.lower() in c['variacao'].lower() or c['variacao'].lower() in var_name.lower()
        ):
            alvo = c
            break

    item = dict(alvo)
    if 'fator_custo' in mapping:
        item['custo'] = round(item['custo'] * mapping['fator_custo'], 2)
    if 'extra_custo' in mapping:
        item['custo'] = round(item['custo'] + mapping['extra_custo'], 2)
    return item


def encontrar_match(linha, ctx):
    var_sku = (linha['var_sku'] or '').upper()
    parent_sku = (linha['parent_sku'] or '').upper()

    if var_sku and var_sku in ctx['por_sku']:
        return ctx['por_sku'][var_sku], 'SKU Direto'
    if parent_sku and parent_sku in ctx['por_sku']:
        return ctx['por_sku'][parent_sku], 'SKU Pai'
    if var_sku:
        prefixo = var_sku.split('-')[0]
        if prefixo in ctx['por_sku']:
            return ctx['por_sku'][prefixo], 'Prefixo SKU'

    var_id = linha['var_id']
    if var_id and var_id in ctx['por_var']:
        alvo = ctx['por_var'][var_id].get('loja_sku', '').upper()
        if alvo in ctx['por_sku']:
            return ctx['por_sku'][alvo], 'ID Variante (Mapeado)'

    chave = f"{linha['prod_id']}_{var_id}"
    if chave in ctx['por_prod_var']:
        alvo = ctx['por_prod_var'][chave].get('loja_sku', '').upper()
        if alvo in ctx['por_sku']:
            return ctx['por_sku'][alvo], 'Anúncio/Variante (Mapeado)'

    if linha['prod_name']:
        alvo = ctx['por_nome_anterior'].get(por_nome_key(linha['prod_name']))
        if alvo:
            sku = alvo.get('loja_sku', '').upper()
            if sku in ctx['por_sku']:
                return ctx['por_sku'][sku], 'Título Shopee (Mapeado)'

    melhor = match_similaridade(linha, ctx)
    if melhor:
        return melhor[0], melhor[1]

    exato = aplicar_mapeamento_exato(linha['prod_id'], linha['var_name'], ctx['por_nome'])
    if exato:
        return exato, 'Mapeamento Exato'

    return None, 'Não Encontrado'


def match_similaridade(linha, ctx):
    """Similaridade de titulo apenas com confianca minima (55% de similaridade)."""
    normalizado = por_nome_key(linha['prod_name'])
    palavras = set(re.findall(r'\b\w{3,}\b', normalizado))
    melhor_item = None
    melhor_score = 0
    melhor_razao = 0.0
    for item in ctx['itens']:
        palavras_loja = set(re.findall(r'\b\w{3,}\b', por_nome_key(item['nome'])))
        sobreposicao = len(palavras.intersection(palavras_loja))
        razao = difflib.SequenceMatcher(None, normalizado, por_nome_key(item['nome'])).ratio()
        score = (sobreposicao * 2) + razao
        if score > melhor_score:
            melhor_score = score
            melhor_razao = razao
            melhor_item = item
    if melhor_item and melhor_score >= 4.0 and melhor_razao >= 0.55:
        return melhor_item, f'Similaridade Título ({int(melhor_razao * 100)}%)'
    return None


def calcular_taxa_shopee(preco):
    """Retorna (taxa, faixa) conforme as regras oficiais Shopee 2026."""
    if preco < 9.00:
        return (preco * 0.20) + (preco * 0.50), 'Sub R$ 9: 20% + 50%'
    if preco <= 79.99:
        return (preco * 0.20) + 4.50, 'Até R$ 79,99: 20% + R$ 4,50'
    if preco <= 99.99:
        return (preco * 0.14) + 16.00, 'R$ 80 a 99,99: 14% + R$ 16,00'
    if preco <= 199.99:
        return (preco * 0.14) + 20.00, 'R$ 100 a 199,99: 14% + R$ 20,00'
    return (preco * 0.14) + 26.00, 'Acima de R$ 200: 14% + R$ 26,00'


def desconto_maximo_seguro(preco_ancora, custo):
    """Maior % de desconto que mantem a margem acima de 20% e 10%."""
    max_20 = 0
    max_10 = 0
    for d in range(0, 91):
        preco = preco_ancora * (1.0 - d / 100.0)
        taxa, _ = calcular_taxa_shopee(preco)
        lucro = preco - taxa - custo - EMBALAGEM
        margem = (lucro / preco * 100.0) if preco > 0 else 0.0
        if lucro > 0 and margem >= 20.0:
            max_20 = d
        if lucro > 0 and margem >= 10.0:
            max_10 = d
    return max_20, max_10


def processar():
    loja_data = carregar_json(ARQUIVO_PRODUTOS_JSON)
    if loja_data is None:
        loja_data = carregar_janela(ARQUIVO_PRODUTOS_JS, 'DADOS_PRODUTOS')
    if loja_data is None:
        sys.exit('Base da loja nao encontrada. Execute primeiro scripts/etl/gerar_dados.py')

    caminho_planilha = localizar_planilha()
    if not caminho_planilha:
        sys.exit(f'Nenhuma planilha encontrada em {relativo(RAW_DIR)}')

    itens_loja, por_sku, por_nome = montar_indice_loja(loja_data)
    por_prod_var, por_var, por_nome_anterior = carregar_mapa_anterior()

    ctx = {
        'itens': itens_loja,
        'por_sku': por_sku,
        'por_nome': por_nome,
        'por_prod_var': por_prod_var,
        'por_var': por_var,
        'por_nome_anterior': por_nome_anterior,
    }

    print(f'Planilha de origem: {relativo(caminho_planilha)}')
    wb = carregar_workbook(caminho_planilha)
    linhas = (ler_exportacao_anuncios if eh_exportacao_anuncios(caminho_planilha) else ler_mass_update)(wb.active)
    print(f'Total de linhas para analise: {len(linhas)}')

    itens_resultado = []
    matched_count = 0

    for linha in linhas:
        preco_cad = linha['preco_cad']
        preco_promo = linha['preco_promo']
        match, match_type = encontrar_match(linha, ctx)

        if match:
            matched_count += 1
            custo = match.get('custo', 0.0)
            preco_loja = match.get('preco_venda', 0.0)
            estoque_loja = match.get('estoque', 0)
            loja_sku = match.get('sku', '')
            loja_nome = match.get('nome', '')
            loja_var = match.get('variacao', '')
            img = linha['img'] or match.get('img', '')
        else:
            # Sem vinculo com o catalogo: custo desconhecido (0.0 nao entra nas medias da UI)
            custo = 0.0
            preco_loja = 0.0
            estoque_loja = 0
            loja_sku = ''
            loja_nome = linha['prod_name']
            loja_var = linha['var_name']
            img = linha['img']

        taxa, faixa = calcular_taxa_shopee(preco_promo)
        repasse = preco_promo - taxa
        lucro = repasse - custo - EMBALAGEM
        margem = (lucro / preco_promo * 100.0) if preco_promo > 0 else 0.0
        desconto_real = ((preco_cad - preco_promo) / preco_cad * 100.0) if preco_cad > 0 else 0.0

        taxa_cad, _ = calcular_taxa_shopee(preco_cad)
        lucro_cad = preco_cad - taxa_cad - custo - EMBALAGEM
        margem_cad = (lucro_cad / preco_cad * 100.0) if preco_cad > 0 else 0.0
        desc_max_20, desc_max_10 = desconto_maximo_seguro(preco_cad, custo)

        itens_resultado.append({
            'prod_id': linha['prod_id'],
            'prod_name': linha['prod_name'],
            'var_id': linha['var_id'],
            'var_name': linha['var_name'],
            'parent_sku': linha['parent_sku'],
            'var_sku': linha['var_sku'],
            'shopee_cad_price': round(preco_cad, 2),
            'shopee_promo_price': round(preco_promo, 2),
            'shopee_price': round(preco_promo, 2),
            'shopee_real_discount_pct': round(desconto_real, 1),
            'shopee_stock': int(linha['qtd']),
            'img': img,
            'match_type': match_type,
            'loja_sku': loja_sku,
            'loja_nome': loja_nome,
            'loja_variacao': loja_var,
            'custo': round(custo, 2),
            'preco_loja': round(preco_loja, 2),
            'estoque_loja': estoque_loja,
            'embalagem': EMBALAGEM,
            'tier': faixa,
            'taxa_shopee': round(taxa, 2),
            'repasse_shopee': round(repasse, 2),
            'lucro_real': round(lucro, 2),
            'margem_real': round(margem, 1),
            'lucro_cad': round(lucro_cad, 2),
            'margem_cad': round(margem_cad, 1),
            'desc_max_seguro_20': desc_max_20,
            'desc_max_seguro_10': desc_max_10,
        })

    if not itens_resultado:
        sys.exit('Nenhum item lido da planilha.')

    total = len(itens_resultado)
    media = lambda campo: sum(it[campo] for it in itens_resultado) / total
    saudaveis = sum(1 for it in itens_resultado if it['margem_real'] >= 20.0)
    margem_baixa = sum(1 for it in itens_resultado if 0.0 <= it['margem_real'] < 20.0)
    prejuizo = sum(1 for it in itens_resultado if it['lucro_real'] < 0.0)

    resultado = {
        'total_shopee': total,
        'arquivo_origem': os.path.basename(caminho_planilha),
        'preco_medio_cad': round(media('shopee_cad_price'), 2),
        'preco_medio_promo': round(media('shopee_promo_price'), 2),
        'desconto_medio_praticado': round(media('shopee_real_discount_pct'), 1),
        'custo_medio': round(media('custo'), 2),
        'lucro_medio_real': round(media('lucro_real'), 2),
        'margem_media_real': round(media('margem_real'), 1),
        'qtd_saudavel': saudaveis,
        'qtd_margem_baixa': margem_baixa,
        'qtd_prejuizo': prejuizo,
        'itens': itens_resultado,
    }

    os.makedirs(WEB_DATA_DIR, exist_ok=True)
    os.makedirs(PROC_DIR, exist_ok=True)

    with open(ARQUIVO_AUDITORIA_JS, 'w', encoding='utf-8') as f:
        f.write('window.DADOS_SHOPEE_REAL = ' + json.dumps(resultado, ensure_ascii=False) + ';\n')
    with open(ARQUIVO_AUDITORIA_JSON, 'w', encoding='utf-8') as f:
        json.dump(resultado, f, ensure_ascii=False, indent=2)

    print(f'Match: {matched_count}/{total} itens')
    print(f'Preco Ancora Medio: R$ {resultado["preco_medio_cad"]:.2f}')
    print(f'Preco com Desconto Real: R$ {resultado["preco_medio_promo"]:.2f}')
    print(f'Desconto Real Medio: {resultado["desconto_medio_praticado"]:.1f}% OFF')
    print(f'Lucro Real Medio: R$ {resultado["lucro_medio_real"]:.2f}')
    print(f'Margem Real Media: {resultado["margem_media_real"]:.1f}%')
    print(f'Itens saudaveis / margem baixa / prejuizo: {saudaveis} / {margem_baixa} / {prejuizo}')
    print(f'Arquivos atualizados: {relativo(ARQUIVO_AUDITORIA_JS)} e {relativo(ARQUIVO_AUDITORIA_JSON)}')


if __name__ == '__main__':
    processar()
